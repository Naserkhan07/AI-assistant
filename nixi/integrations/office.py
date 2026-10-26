"""Optional, shell-free Office document skills using Python file libraries."""
from __future__ import annotations

from pathlib import Path
from typing import Any, Iterable

from nixi.core.file_actions import FileActionService


class OfficeIntegration:
    def __init__(self, file_actions: FileActionService):
        self.files = file_actions

    def create_word_document(self, path: str, title: str, body: str = "") -> str:
        try:
            from docx import Document
        except ImportError:
            return "Word document support is unavailable; install the python-docx dependency."
        try:
            target = self.files._resolve(path)
            if target.exists():
                return f"File already exists; nothing overwritten: {target}"
            if target.suffix.lower() != ".docx" or not target.parent.is_dir():
                return "Provide a .docx path in an existing folder under your home directory."
            doc = Document()
            doc.add_heading(title[:200], level=1)
            if body:
                for paragraph in body[:50000].splitlines()[:200]:
                    doc.add_paragraph(paragraph[:2000])
            doc.save(str(target))
            self.files.undo.record_create(target)
            check = Document(str(target))
            if not check.paragraphs or check.paragraphs[0].text != title[:200]:
                return "Document was saved but verification failed."
            return f"Created and verified Word document: {target}"
        except Exception as exc:
            return f"Could not create Word document: {exc}"

    def read_word_document(self, path: str) -> str:
        try:
            from docx import Document
            target = self.files._resolve(path)
            if not target.is_file() or target.suffix.lower() != ".docx":
                return f"Word document not found: {target}"
            text = "\n".join(p.text for p in Document(str(target)).paragraphs)
            return text[:12000] or "The document contains no paragraph text."
        except ImportError:
            return "Word document support is unavailable; install python-docx."
        except Exception as exc:
            return f"Could not read Word document: {exc}"

    def create_workbook(self, path: str, sheet_name: str = "Sheet1", rows: list[list[Any]] | None = None) -> str:
        try:
            from openpyxl import Workbook, load_workbook
        except ImportError:
            return "Excel workbook support is unavailable; install the openpyxl dependency."
        try:
            target = self.files._resolve(path)
            if target.exists():
                return f"File already exists; nothing overwritten: {target}"
            if target.suffix.lower() != ".xlsx" or not target.parent.is_dir():
                return "Provide an .xlsx path in an existing folder under your home directory."
            wb = Workbook()
            ws = wb.active
            ws.title = (sheet_name or "Sheet1")[:31]
            safe_rows = (rows or [])[:100]
            for row in safe_rows:
                values = []
                for value in row[:50]:
                    if value is None:
                        values.append(None)
                    elif isinstance(value, str):
                        text = value[:2000]
                        values.append("'" + text if text.lstrip().startswith(("=", "+", "-", "@")) else text)
                    elif isinstance(value, (int, float, bool)):
                        values.append(value)
                    else:
                        values.append(str(value)[:2000])
                ws.append(values)
            wb.save(target)
            self.files.undo.record_create(target)
            verify = load_workbook(target, read_only=True, data_only=True)
            sheet = verify[ws.title]
            ok = sheet.max_row == max(1, len(safe_rows))
            verify.close()
            return f"Created and verified Excel workbook: {target}" if ok else "Workbook saved but verification failed."
        except Exception as exc:
            return f"Could not create Excel workbook: {exc}"

    def read_workbook(self, path: str, sheet_name: str = "") -> list[list[Any]] | str:
        try:
            from openpyxl import load_workbook
        except ImportError:
            return "Excel workbook support is unavailable; install openpyxl."
        try:
            target = self.files._resolve(path)
            if not target.is_file() or target.suffix.lower() != ".xlsx":
                return f"Workbook not found: {target}"
            wb = load_workbook(target, read_only=True, data_only=True)
            ws = wb[sheet_name] if sheet_name and sheet_name in wb.sheetnames else wb.active
            rows = [[cell for cell in row] for row in ws.iter_rows(values_only=True, max_row=100, max_col=50)]
            wb.close()
            return rows
        except Exception as exc:
            return f"Could not read workbook: {exc}"
